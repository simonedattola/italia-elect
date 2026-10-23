#!/usr/bin/env python3
"""Scarica gli zip open data Eligendo e scrive scrutini compatti.

Fonte: catalogo pubblicato in
https://elezionistorico.interno.gov.it/eligendo/opendata.php
I file stanno su https://dait.interno.gov.it/documenti/opendata/{cartella}/{nome}.

Non ricostruisce celle mancanti. I file di preferenza non diventano voti di lista.
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import urllib.request
import zipfile
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / "eligendo"
OUT = ROOT / "src" / "data" / "generated"
RETRIEVED = "2026-10-23"
BASE = "https://dait.interno.gov.it/documenti/opendata/"
CATALOG_PAGE = "https://elezionistorico.interno.gov.it/eligendo/opendata.php"
UA = "ItaliaElectIngest/1.0 (research; +https://github.com/simonedattola/italia-elect)"

REGIONS = {
    "PIEMONTE": "Piemonte",
    "VALLE D'AOSTA": "Valle d'Aosta",
    "VALLE D AOSTA": "Valle d'Aosta",
    "VALLE DAOSTA": "Valle d'Aosta",
    "LOMBARDIA": "Lombardia",
    "TRENTINO-ALTO ADIGE": "Trentino-Alto Adige",
    "TRENTINO ALTO ADIGE": "Trentino-Alto Adige",
    "TRENTINO-ALTO ADIGE/SUDTIROL": "Trentino-Alto Adige",
    "VENETO": "Veneto",
    "FRIULI-VENEZIA GIULIA": "Friuli-Venezia Giulia",
    "FRIULI VENEZIA GIULIA": "Friuli-Venezia Giulia",
    "LIGURIA": "Liguria",
    "EMILIA-ROMAGNA": "Emilia-Romagna",
    "EMILIA ROMAGNA": "Emilia-Romagna",
    "TOSCANA": "Toscana",
    "UMBRIA": "Umbria",
    "MARCHE": "Marche",
    "LAZIO": "Lazio",
    "ABRUZZO": "Abruzzo",
    "MOLISE": "Molise",
    "CAMPANIA": "Campania",
    "PUGLIA": "Puglia",
    "BASILICATA": "Basilicata",
    "CALABRIA": "Calabria",
    "SICILIA": "Sicilia",
    "SARDEGNA": "Sardegna",
    "PROVINCIA AUTONOMA DI TRENTO": "Trentino-Alto Adige",
    "PROVINCIA AUTONOMA DI BOLZANO": "Trentino-Alto Adige",
    "BOLZANO": "Trentino-Alto Adige",
    "TRENTO": "Trentino-Alto Adige",
}

# Capoluoghi di provincia e città metropolitane, più i tre poli di Barletta-Andria-Trani
# e Carbonia (Sud Sardegna). La chiave è il nome normalizzato nel file.
CAPOLUOGHI = {
    "AGRIGENTO": "Agrigento",
    "ALESSANDRIA": "Alessandria",
    "ANCONA": "Ancona",
    "AOSTA": "Aosta",
    "AREZZO": "Arezzo",
    "ASCOLI PICENO": "Ascoli Piceno",
    "ASTI": "Asti",
    "AVELLINO": "Avellino",
    "BARI": "Bari",
    "BARLETTA": "Barletta",
    "ANDRIA": "Andria",
    "TRANI": "Trani",
    "BELLUNO": "Belluno",
    "BENEVENTO": "Benevento",
    "BERGAMO": "Bergamo",
    "BIELLA": "Biella",
    "BOLOGNA": "Bologna",
    "BOLZANO": "Bolzano",
    "BRESCIA": "Brescia",
    "BRINDISI": "Brindisi",
    "CAGLIARI": "Cagliari",
    "CALTANISSETTA": "Caltanissetta",
    "CAMPOBASSO": "Campobasso",
    "CARBONIA": "Carbonia",
    "CASERTA": "Caserta",
    "CATANIA": "Catania",
    "CATANZARO": "Catanzaro",
    "CHIETI": "Chieti",
    "COMO": "Como",
    "COSENZA": "Cosenza",
    "CREMONA": "Cremona",
    "CROTONE": "Crotone",
    "CUNEO": "Cuneo",
    "ENNA": "Enna",
    "FERMO": "Fermo",
    "FERRARA": "Ferrara",
    "FIRENZE": "Firenze",
    "FOGGIA": "Foggia",
    "FORLI": "Forlì",
    "FROSINONE": "Frosinone",
    "GENOVA": "Genova",
    "GORIZIA": "Gorizia",
    "GROSSETO": "Grosseto",
    "IMPERIA": "Imperia",
    "ISERNIA": "Isernia",
    "L AQUILA": "L'Aquila",
    "LAQUILA": "L'Aquila",
    "LA SPEZIA": "La Spezia",
    "LATINA": "Latina",
    "LECCE": "Lecce",
    "LECCO": "Lecco",
    "LIVORNO": "Livorno",
    "LODI": "Lodi",
    "LUCCA": "Lucca",
    "MACERATA": "Macerata",
    "MANTOVA": "Mantova",
    "MASSA": "Massa",
    "MATERA": "Matera",
    "MESSINA": "Messina",
    "MILANO": "Milano",
    "MODENA": "Modena",
    "MONZA": "Monza",
    "NAPOLI": "Napoli",
    "NOVARA": "Novara",
    "NUORO": "Nuoro",
    "ORISTANO": "Oristano",
    "PADOVA": "Padova",
    "PALERMO": "Palermo",
    "PARMA": "Parma",
    "PAVIA": "Pavia",
    "PERUGIA": "Perugia",
    "PESARO": "Pesaro",
    "PESCARA": "Pescara",
    "PIACENZA": "Piacenza",
    "PISA": "Pisa",
    "PISTOIA": "Pistoia",
    "PORDENONE": "Pordenone",
    "POTENZA": "Potenza",
    "PRATO": "Prato",
    "RAGUSA": "Ragusa",
    "RAVENNA": "Ravenna",
    "REGGIO CALABRIA": "Reggio Calabria",
    "REGGIO DI CALABRIA": "Reggio Calabria",
    "REGGIO NELL EMILIA": "Reggio Emilia",
    "REGGIO EMILIA": "Reggio Emilia",
    "RIETI": "Rieti",
    "RIMINI": "Rimini",
    "ROMA": "Roma",
    "ROMA CAPITALE": "Roma",
    "ROVIGO": "Rovigo",
    "SALERNO": "Salerno",
    "SASSARI": "Sassari",
    "SAVONA": "Savona",
    "SIENA": "Siena",
    "SIRACUSA": "Siracusa",
    "SONDRIO": "Sondrio",
    "TARANTO": "Taranto",
    "TERAMO": "Teramo",
    "TERNI": "Terni",
    "TORINO": "Torino",
    "TRAPANI": "Trapani",
    "TRENTO": "Trento",
    "TREVISO": "Treviso",
    "TRIESTE": "Trieste",
    "UDINE": "Udine",
    "VARESE": "Varese",
    "VENEZIA": "Venezia",
    "VERBANIA": "Verbania",
    "VERCELLI": "Vercelli",
    "VERONA": "Verona",
    "VIBO VALENTIA": "Vibo Valentia",
    "VICENZA": "Vicenza",
    "VITERBO": "Viterbo",
}

PROVINCE_REGION = {
    "TORINO": "Piemonte",
    "VERCELLI": "Piemonte",
    "NOVARA": "Piemonte",
    "CUNEO": "Piemonte",
    "ASTI": "Piemonte",
    "ALESSANDRIA": "Piemonte",
    "BIELLA": "Piemonte",
    "VERBANO-CUSIO-OSSOLA": "Piemonte",
    "VERBANO CUSIO OSSOLA": "Piemonte",
    "AOSTA": "Valle d'Aosta",
    "VARESE": "Lombardia",
    "COMO": "Lombardia",
    "SONDRIO": "Lombardia",
    "MILANO": "Lombardia",
    "BERGAMO": "Lombardia",
    "BRESCIA": "Lombardia",
    "PAVIA": "Lombardia",
    "CREMONA": "Lombardia",
    "MANTOVA": "Lombardia",
    "LECCO": "Lombardia",
    "LODI": "Lombardia",
    "MONZA": "Lombardia",
    "MONZA E DELLA BRIANZA": "Lombardia",
    "BOLZANO": "Trentino-Alto Adige",
    "TRENTO": "Trentino-Alto Adige",
    "VERONA": "Veneto",
    "VICENZA": "Veneto",
    "BELLUNO": "Veneto",
    "TREVISO": "Veneto",
    "VENEZIA": "Veneto",
    "PADOVA": "Veneto",
    "ROVIGO": "Veneto",
    "UDINE": "Friuli-Venezia Giulia",
    "GORIZIA": "Friuli-Venezia Giulia",
    "TRIESTE": "Friuli-Venezia Giulia",
    "PORDENONE": "Friuli-Venezia Giulia",
    "IMPERIA": "Liguria",
    "SAVONA": "Liguria",
    "GENOVA": "Liguria",
    "LA SPEZIA": "Liguria",
    "PIACENZA": "Emilia-Romagna",
    "PARMA": "Emilia-Romagna",
    "REGGIO EMILIA": "Emilia-Romagna",
    "REGGIO NELL EMILIA": "Emilia-Romagna",
    "MODENA": "Emilia-Romagna",
    "BOLOGNA": "Emilia-Romagna",
    "FERRARA": "Emilia-Romagna",
    "RAVENNA": "Emilia-Romagna",
    "FORLI": "Emilia-Romagna",
    "FORLI-CESENA": "Emilia-Romagna",
    "RIMINI": "Emilia-Romagna",
    "MASSA": "Toscana",
    "MASSA-CARRARA": "Toscana",
    "LUCCA": "Toscana",
    "PISTOIA": "Toscana",
    "FIRENZE": "Toscana",
    "LIVORNO": "Toscana",
    "PISA": "Toscana",
    "AREZZO": "Toscana",
    "SIENA": "Toscana",
    "GROSSETO": "Toscana",
    "PRATO": "Toscana",
    "PERUGIA": "Umbria",
    "TERNI": "Umbria",
    "PESARO": "Marche",
    "PESARO E URBINO": "Marche",
    "ANCONA": "Marche",
    "MACERATA": "Marche",
    "ASCOLI PICENO": "Marche",
    "FERMO": "Marche",
    "VITERBO": "Lazio",
    "RIETI": "Lazio",
    "ROMA": "Lazio",
    "LATINA": "Lazio",
    "FROSINONE": "Lazio",
    "L AQUILA": "Abruzzo",
    "TERAMO": "Abruzzo",
    "PESCARA": "Abruzzo",
    "CHIETI": "Abruzzo",
    "CAMPOBASSO": "Molise",
    "ISERNIA": "Molise",
    "CASERTA": "Campania",
    "BENEVENTO": "Campania",
    "NAPOLI": "Campania",
    "AVELLINO": "Campania",
    "SALERNO": "Campania",
    "FOGGIA": "Puglia",
    "BARI": "Puglia",
    "TARANTO": "Puglia",
    "BRINDISI": "Puglia",
    "LECCE": "Puglia",
    "BARLETTA": "Puglia",
    "BARLETTA-ANDRIA-TRANI": "Puglia",
    "POTENZA": "Basilicata",
    "MATERA": "Basilicata",
    "COSENZA": "Calabria",
    "CATANZARO": "Calabria",
    "REGGIO CALABRIA": "Calabria",
    "REGGIO DI CALABRIA": "Calabria",
    "CROTONE": "Calabria",
    "VIBO VALENTIA": "Calabria",
    "TRAPANI": "Sicilia",
    "PALERMO": "Sicilia",
    "MESSINA": "Sicilia",
    "AGRIGENTO": "Sicilia",
    "CALTANISSETTA": "Sicilia",
    "ENNA": "Sicilia",
    "CATANIA": "Sicilia",
    "RAGUSA": "Sicilia",
    "SIRACUSA": "Sicilia",
    "SASSARI": "Sardegna",
    "NUORO": "Sardegna",
    "CAGLIARI": "Sardegna",
    "ORISTANO": "Sardegna",
    "SUD SARDEGNA": "Sardegna",
    "CARBONIA": "Sardegna",
    "CARBONIA-IGLESIAS": "Sardegna",
    "MEDIO CAMPIDANO": "Sardegna",
    "OGLIASTRA": "Sardegna",
    "OLBIA-TEMPIO": "Sardegna",
}


def fold(value: str) -> str:
    text = unicodedata.normalize("NFD", value or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper().replace("’", "'").replace("`", "'").replace("'", " ")
    text = text.replace("-", " ")
    text = re.sub(r"[^A-Z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: str) -> str:
    text = fold(value).lower().replace(" ", "-")
    text = re.sub(r"-+", "-", text).strip("-")
    return text[:96] or "lista"


def header_key(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", fold(value))


def parse_number(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip().strip('"').replace(" ", "").replace("%", "")
    if text in {"", "-", "–", "—"}:
        return None
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d+,\d+", text):
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def round2(value: float) -> float:
    return round(value + 1e-9, 2)


def fetch(url: str, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=180) as response:
            dest.write_bytes(response.read())
    except Exception:
        if dest.exists():
            dest.unlink()
        raise
    return dest


def catalog_files() -> list[tuple[str, str]]:
    req = urllib.request.Request(CATALOG_PAGE, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as response:
        html = response.read().decode("utf-8", "replace")
    rows = re.findall(
        r'\[\s*"([^"]+)"\s*,\s*"([^"]*)"\s*,\s*"([^"]+)"\s*,\s*"([^"]+)"',
        html,
    )
    found: list[tuple[str, str]] = []
    for _kind, _year, path, _name in rows:
        if path.startswith(
            ("senato/", "europee/", "regionali/", "comunali/", "referendum/", "camera/camera-20180304", "camera/camera-20220925")
        ):
            found.append((path, path.split("/")[-1]))
    # unique
    uniq = []
    seen = set()
    for path, name in found:
        if path in seen:
            continue
        seen.add(path)
        uniq.append((path, name))
    return uniq


def decode_member(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", "replace")


HEADER_ALIAS = {
    "DESCRCOMUNE": "COMUNE",
    "DESCRIZIONECOMUNE": "COMUNE",
    "DESCRPROV": "PROVINCIA",
    "DESCRREG": "REGIONE",
    "DESCRCIRC": "CIRCOSCRIZIONE",
    "CIRCREG": "REGIONE",
    "CIRCOSCR": "CIRCOSCRIZIONE",
    "DESCRLISTA": "LISTA",
    "QUESITOREFERENDUM": "QUESITO",
    "NUMQUESITO": "NUMREFERENDUM",
    "VOTIVALIDISI": "NUMVOTISI",
    "VOTISI": "NUMVOTISI",
    "VOTINO": "NUMVOTINO",
    "ELETTORITOT": "ELETTORI",
    "ELETTORITOTALI": "ELETTORI",
    "NUMELETTORI": "ELETTORI",
    "NUMELETTORITOTALI": "ELETTORI",
    "ELETTORITOTALE": "ELETTORI",
    "VOTANTITOT": "VOTANTI",
    "VOTANTITOTALI": "VOTANTI",
    "NUMVOTANTITOTALI": "VOTANTI",
    "SKBIANCHE": "BIANCHE",
    "SCHEDEBIANCHE": "BIANCHE",
    "VOTILISTA": "VOTI",
    "TOTVOTI": "VOTI",
    "SEGGILISTA": "SEGGI",
}


def canon_row(raw: dict[str, str]) -> dict[str, str]:
    cleaned: dict[str, str] = {}
    for key, value in raw.items():
        if key is None:
            continue
        token = header_key(str(key))
        text = "" if value is None else str(value).strip().strip('"')
        cleaned[token] = text
        alias = HEADER_ALIAS.get(token)
        if alias and alias not in cleaned:
            cleaned[alias] = text
    return cleaned


def read_table(text: str) -> list[dict[str, str]]:
    sample = text[:4000]
    delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    first = text.splitlines()[0] if text else ""
    probe = [header_key(cell) for cell in next(csv.reader(io.StringIO(first), delimiter=delimiter), [])]
    known = {"LISTA", "DESCRLISTA", "QUESITO", "QUESITOREFERENDUM", "NUMVOTISI", "ELETTORI", "COMUNE", "REGIONE", "DESCRREG", "DATAELEZIONE"}
    if probe and not any(cell in known or cell in HEADER_ALIAS for cell in probe):
        return positional_referendum(text, delimiter)
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        return []
    rows = []
    for row in reader:
        cleaned = canon_row(row)
        if any(cleaned.values()):
            rows.append(cleaned)
    return rows


def positional_referendum(text: str, delimiter: str) -> list[dict[str, str]]:
    """Alcuni zip 1980 non hanno intestazione: regione, provincia, comune, numero, quesito, elettori, maschi, votanti, maschi, sì, no, bianche."""
    rows = []
    for cells in csv.reader(io.StringIO(text), delimiter=delimiter):
        if len(cells) < 11:
            continue
        rows.append(
            {
                "REGIONE": cells[0].strip().strip('"'),
                "PROVINCIA": cells[1].strip().strip('"'),
                "COMUNE": cells[2].strip().strip('"'),
                "NUMREFERENDUM": cells[3].strip().strip('"'),
                "QUESITO": cells[4].strip().strip('"'),
                "ELETTORI": cells[5].strip().strip('"'),
                "VOTANTI": cells[7].strip().strip('"'),
                "NUMVOTISI": cells[9].strip().strip('"'),
                "NUMVOTINO": cells[10].strip().strip('"'),
                "BIANCHE": cells[11].strip().strip('"') if len(cells) > 11 else "",
            }
        )
    return rows


def xlsx_table(blob: bytes) -> list[dict[str, str]]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    iterator = sheet.iter_rows(values_only=True)
    header = next(iterator, None)
    if not header:
        book.close()
        return []
    rows = []
    for values in iterator:
        raw = {str(header[i]): "" if values[i] is None else str(values[i]) for i in range(min(len(header), len(values))) if header[i]}
        cleaned = canon_row(raw)
        if any(cleaned.values()):
            rows.append(cleaned)
    book.close()
    return rows


def zip_tables(path: Path) -> list[tuple[str, list[dict[str, str]]]]:
    tables = []
    with zipfile.ZipFile(path) as archive:
        names = [info.filename for info in archive.infolist() if not info.is_dir()]
        has_comune = any("comune" in name.lower() and "sez" not in name.lower() for name in names)
        for info in archive.infolist():
            if info.is_dir():
                continue
            name = info.filename
            lower = name.lower()
            if "pref" in lower or "estero" in lower:
                continue
            if has_comune and "sez" in lower:
                continue
            if lower.endswith(".xlsx"):
                with archive.open(info) as handle:
                    tables.append((name, xlsx_table(handle.read())))
                continue
            if lower.endswith((".pdf", ".xls", ".png")):
                continue
            with archive.open(info) as handle:
                text = decode_member(handle.read())
            tables.append((name, read_table(text)))
    return tables


def pick(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        if row.get(key):
            return row[key]
    return ""


def pick_num(row: dict[str, str], *keys: str) -> float | None:
    for key in keys:
        if key in row:
            value = parse_number(row[key])
            if value is not None:
                return value
    return None


def region_of(row: dict[str, str]) -> str | None:
    for key in ("REGIONE", "REG", "DENOMINAZIONE"):
        label = fold(row.get(key, ""))
        if not label or label.isdigit():
            continue
        if label in {"ITALIA", "ITALIA ITALIA"}:
            return "Italia"
        mapped = REGIONS.get(label)
        if mapped:
            return mapped
        # "PIEMONTE 1"
        mapped = REGIONS.get(re.sub(r"\s+\d+$", "", label))
        if mapped:
            return mapped
    province = fold(row.get("PROVINCIA") or row.get("PROV") or "")
    if province in PROVINCE_REGION:
        return PROVINCE_REGION[province]
    circ = fold(row.get("CIRCOSCRIZIONE") or row.get("CIRC") or "")
    if circ in REGIONS:
        return REGIONS[circ]
    return None


def is_preference_name(name: str) -> bool:
    # I file «pref» ripetono le preferenze. «Liste&Candidati» invece ha VOTILISTA.
    return "pref" in name.lower()


def classify(name: str, rows: list[dict[str, str]]) -> str:
    lower = name.lower()
    if is_preference_name(name):
        return "pref"
    if "estero" in lower:
        return "estero"
    header = set(rows[0].keys()) if rows else set()
    has_list = bool(header & {"LISTA", "DESCRLISTA"})
    has_si = bool(header & {"NUMVOTISI", "VOTISI", "SI"})
    if has_si and not has_list:
        return "referendum"
    if "scrutin" in lower and not has_list:
        return "turnout"
    if any(token in lower for token in ("liv_naz", "livnaz", "_naz")) and has_list:
        return "lists-naz"
    if any(token in lower for token in ("liv_reg", "livreg")) and has_list:
        return "lists-reg"
    if "comune" in lower or "COMUNE" in header:
        return "comune" if has_list or has_si else "other"
    if has_list:
        return "lists"
    return "other"


def list_name(row: dict[str, str]) -> str:
    return pick(row, "DESCRLISTA", "LISTA").strip()


def list_votes(row: dict[str, str]) -> float | None:
    return pick_num(row, "VOTI", "VOTILISTA", "TOTVOTI")


def list_seats(row: dict[str, str]) -> float | None:
    return pick_num(row, "SEGGILISTA", "SEGGI")


def place_key(row: dict[str, str]) -> str:
    return "|".join(
        [
            fold(pick(row, "COMUNE", "DESCRIZIONECOMUNE")),
            fold(pick(row, "PROVINCIA", "PROV")),
            fold(pick(row, "REGIONE", "REG", "CIRCOSCRIZIONE")),
        ]
    )


def aggregate_lists(rows: list[dict[str, str]]) -> tuple[dict[str, dict[str, float]], dict[str, dict[str, float]], dict[str, dict[str, float]]]:
    """Ritorna voti, seggi e meta (elettori/votanti/bianche) per territorio."""
    votes: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    seats: dict[str, dict[str, float]] = defaultdict(dict)
    meta: dict[str, dict[str, float]] = defaultdict(lambda: {"elettori": 0.0, "votanti": 0.0, "bianche": 0.0})
    seen_list: set[tuple[str, str]] = set()
    seen_place: set[str] = set()
    for row in rows:
        name = list_name(row)
        if not name:
            continue
        territory = region_of(row) or "Italia"
        if territory == "Italia" and pick(row, "COMUNE", "DESCRIZIONECOMUNE"):
            territory = region_of(row) or "Non mappata"
        key_place = place_key(row) or territory
        dedupe = (key_place, fold(name))
        if dedupe in seen_list:
            continue
        seen_list.add(dedupe)
        amount = list_votes(row)
        if amount is None:
            continue
        votes[territory][name] += amount
        seat = list_seats(row)
        if seat is not None and name not in seats[territory]:
            seats[territory][name] = seat
        if key_place not in seen_place:
            seen_place.add(key_place)
            elettori = pick_num(row, "ELETTORI", "ELETTORITOT", "NUMELETTORITOTALI", "ELETTORITOTALI")
            votanti = pick_num(row, "VOTANTI", "VOTANTITOT", "NUMVOTANTITOTALI")
            bianche = pick_num(row, "SKBIANCHE", "SCHEDEBIANCHE", "SCHEDEBIANCHE")
            if elettori:
                meta[territory]["elettori"] += elettori
            if votanti:
                meta[territory]["votanti"] += votanti
            if bianche:
                meta[territory]["bianche"] += bianche
    return votes, seats, meta


def apply_turnout(meta: dict[str, dict[str, float]], rows: list[dict[str, str]]) -> None:
    for row in rows:
        territory = region_of(row) or "Italia"
        elettori = pick_num(row, "ELETTORI", "ELETTORITOT", "NUMELETTORITOTALI", "ELETTORITOTALI")
        votanti = pick_num(row, "VOTANTI", "VOTANTITOT", "NUMVOTANTITOTALI")
        bianche = pick_num(row, "SKBIANCHE", "SCHEDEBIANCHE")
        if elettori:
            meta[territory]["elettori"] = elettori
        if votanti:
            meta[territory]["votanti"] = votanti
        if bianche:
            meta[territory]["bianche"] = bianche


def lists_from(votes: dict[str, float], seats: dict[str, float], scope: str) -> list[dict]:
    total = sum(votes.values()) or 1
    used: dict[str, int] = {}
    rows = []
    for name, amount in sorted(votes.items(), key=lambda item: -item[1]):
        base = slugify(name)
        n = used.get(base, 0) + 1
        used[base] = n
        ident = base if n == 1 else f"{base}-{n}"
        seat = seats.get(name)
        rows.append(
            {
                "id": ident,
                "name": name,
                "votes": int(round(amount)),
                "percentage": round2(100 * amount / total),
                "seats": int(seat) if seat is not None else None,
                "seatsScope": scope if seat is not None else "non-disponibile",
            }
        )
    return rows


def territory_block(name: str, votes: dict[str, float], seats: dict[str, float], meta: dict[str, float], scope: str) -> dict:
    elettori = meta.get("elettori") or None
    votanti = meta.get("votanti") or None
    bianche = meta.get("bianche") or None
    valid = sum(votes.values())
    turnout = round2(100 * votanti / elettori) if elettori and votanti else None
    return {
        "id": slugify(name),
        "name": name,
        "electorate": int(elettori) if elettori else None,
        "voters": int(votanti) if votanti else None,
        "blankBallots": int(bianche) if bianche else None,
        "validVotes": int(round(valid)) if valid else None,
        "turnout": turnout,
        "lists": lists_from(votes, seats, scope),
    }


def source(url: str, notes: str) -> dict:
    return {
        "provider": "Ministero dell'Interno — DAIT, open data Eligendo",
        "url": url,
        "retrievedAt": RETRIEVED,
        "notes": notes,
    }


def summary_of(record: dict) -> dict:
    lists = record["lists"][:6]
    return {
        "id": record["id"],
        "level": record["level"],
        "chamber": record["chamber"],
        "date": record["date"],
        "year": record["year"],
        "title": record["title"],
        "territory": record["territory"],
        "turnout": record["turnout"],
        "listCount": len(record["lists"]) if record["level"] != "comunali" else len(record.get("municipalities") or record["lists"]),
        "hasVotes": any(item.get("votes") for item in record["lists"]),
        "hasGeography": bool(record.get("regions")),
        "topLists": [{"id": item["id"], "name": item["name"], "percentage": item["percentage"]} for item in lists],
        "coverageNotes": record["coverageNotes"],
        "source": record["source"],
    }


def date_from_name(name: str) -> str | None:
    match = re.search(r"(19|20)\d{6}", name)
    if not match:
        return None
    raw = match.group(0)
    return f"{raw[:4]}-{raw[4:6]}-{raw[6:8]}"


def build_record(family: str, path: str, tables: list[tuple[str, list[dict[str, str]]]]) -> dict | None:
    date = date_from_name(path)
    if not date:
        return None
    year = int(date[:4])
    useful = [(name, rows) for name, rows in tables if rows]
    if not useful and family != "europee":
        return None
    kinds = {name: classify(name, rows) for name, rows in useful}
    url = BASE + path
    if family == "referendum":
        return build_referendum(date, year, url, useful)
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for name, rows in useful:
        grouped[kinds[name]].extend(rows)
    if grouped["lists-naz"]:
        list_rows = grouped["lists-naz"]
        regional_rows = grouped["lists-reg"]
    elif grouped["lists-reg"] and not grouped["comune"]:
        list_rows = grouped["lists-reg"]
        regional_rows = []
    else:
        list_rows = grouped["comune"] or grouped["lists"] or grouped["lists-reg"]
        regional_rows = []
    turnout_rows = grouped["turnout"]
    if not list_rows:
        return None
    # Se il file nazionale esiste, non sommare anche i comuni.
    votes, seats, meta = aggregate_lists(list_rows)
    if regional_rows and "Italia" in votes:
        reg_votes, reg_seats, reg_meta = aggregate_lists(regional_rows)
        for territory, bucket in reg_votes.items():
            if territory != "Italia":
                votes[territory] = bucket
                seats[territory] = reg_seats.get(territory, {})
                if territory in reg_meta:
                    meta[territory] = reg_meta[territory]
    if turnout_rows:
        apply_turnout(meta, turnout_rows)
    italy_votes = votes.get("Italia")
    if not italy_votes:
        italy_votes = defaultdict(float)
        for territory, bucket in votes.items():
            if territory == "Italia":
                continue
            for name, amount in bucket.items():
                italy_votes[name] += amount
        if not italy_votes:
            return None
    italy_seats = seats.get("Italia", {})
    if not italy_seats:
        # i seggi regionali non si sommano se il file nazionale li ha già; se mancano, somma
        summed: dict[str, float] = defaultdict(float)
        for territory, bucket in seats.items():
            if territory == "Italia":
                continue
            for name, amount in bucket.items():
                summed[name] += amount
        italy_seats = summed
    scope = "proporzionale" if year >= 2017 and family in {"senato", "camera"} else ("totale" if italy_seats else "non-disponibile")
    italy_meta = meta.get("Italia", {"elettori": 0, "votanti": 0, "bianche": 0})
    if not italy_meta.get("elettori"):
        italy_meta = {
            "elettori": sum(item.get("elettori", 0) for key, item in meta.items() if key not in {"Italia"}),
            "votanti": sum(item.get("votanti", 0) for key, item in meta.items() if key not in {"Italia"}),
            "bianche": sum(item.get("bianche", 0) for key, item in meta.items() if key not in {"Italia"}),
        }
    national = territory_block("Italia", italy_votes, italy_seats, italy_meta, scope)
    regions = []
    for territory, bucket in sorted(votes.items()):
        if territory in {"Italia", "Non mappata"}:
            continue
        regions.append(territory_block(territory, bucket, seats.get(territory, {}), meta.get(territory, {}), scope))
    titles = {
        "senato": f"Politiche {year} — Senato",
        "camera": f"Politiche {year} — Camera, totale di lista",
        "europee": f"Elezioni europee {year} — Italia",
        "regionali": f"Regionali {date}",
    }
    levels = {"senato": "nazionali", "camera": "nazionali", "europee": "europee", "regionali": "regionali"}
    chambers = {"senato": "senato", "camera": "camera", "europee": "parlamento-europeo", "regionali": "consiglio"}
    ident = {
        "senato": f"politiche-senato-{year}",
        "camera": f"politiche-camera-{year}-eligendo",
        "europee": f"europee-{year}",
        "regionali": f"regionali-ciclo-{date}",
    }[family]
    notes = [
        f"Totali ricavati da {path}. Le righe ripetute per lo stesso comune e la stessa lista sono contate una volta.",
        "I file di preferenza non entrano nei voti. Dove la colonna seggi manca, il seggio resta vuoto.",
    ]
    if scope == "proporzionale":
        notes.append("Dal 2017 la colonna seggi di lista del file è la quota proporzionale, non i seggi uninominali.")
    return {
        "id": ident,
        "level": levels[family],
        "chamber": chambers[family],
        "date": date,
        "year": year,
        "title": titles[family],
        "territory": "Italia",
        "source": source(url, f"Zip open data {path}."),
        "coverageNotes": notes,
        "turnout": national["turnout"],
        "electorate": national["electorate"],
        "voters": national["voters"],
        "blankBallots": national["blankBallots"],
        "validVotes": national["validVotes"],
        "lists": national["lists"],
        "regions": regions,
        "colleges": [],
    }


def build_referendum(date: str, year: int, url: str, tables: list[tuple[str, list[dict[str, str]]]]) -> dict | None:
    rows = []
    for _name, table in tables:
        if table and ("NUMVOTISI" in table[0] or "VOTISI" in table[0] or "SI" in table[0]):
            rows.extend(table)
    if not rows:
        return None
    # quesito -> territory -> sums, dedupe comune
    bucket: dict[str, dict[str, dict[str, float]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    seen: set[tuple[str, str]] = set()
    titles: dict[str, str] = {}
    for row in rows:
        qid = pick(row, "NUMREFERENDUM", "NUMQUESITO") or "1"
        title = pick(row, "QUESITO", "DESCRIZIONE") or f"Quesito {qid}"
        titles[qid] = title
        territory = region_of(row) or "Non mappata"
        place = place_key(row) or territory
        marker = (qid, place)
        if marker in seen:
            continue
        seen.add(marker)
        si = pick_num(row, "NUMVOTISI", "VOTISI", "SI") or 0
        no = pick_num(row, "NUMVOTINO", "VOTINO", "NO")
        if no is None:
            valid = pick_num(row, "VOTIVALIDI")
            no = max(0, valid - si) if valid is not None else 0
        elettori = pick_num(row, "ELETTORI", "ELETTORITOT", "NUMELETTORITOTALI") or 0
        votanti = pick_num(row, "VOTANTI", "VOTANTITOT", "NUMVOTANTITOTALI") or 0
        bianche = pick_num(row, "SKBIANCHE", "SCHEDEBIANCHE") or 0
        target = bucket[qid][territory]
        target["si"] += si
        target["no"] += no
        target["elettori"] += elettori
        target["votanti"] += votanti
        target["bianche"] += bianche
    questions = []
    for qid, territories in bucket.items():
        national = {"si": 0.0, "no": 0.0, "elettori": 0.0, "votanti": 0.0, "bianche": 0.0}
        regions = []
        for territory, values in territories.items():
            if territory == "Non mappata":
                for key in national:
                    national[key] += values[key]
                continue
            for key in national:
                national[key] += values[key]
            regions.append(question_territory(territory, values))
        questions.append(
            {
                "id": f"referendum-{date}-q{qid}",
                "number": qid,
                "title": titles.get(qid, f"Quesito {qid}"),
                "national": question_territory("Italia", national),
                "regions": regions,
            }
        )
    if not questions:
        return None
    # Un record archivio per quesito, più il pacchetto compatto.
    return {"date": date, "year": year, "url": url, "questions": questions}


def question_territory(name: str, values: dict[str, float]) -> dict:
    si = values.get("si", 0)
    no = values.get("no", 0)
    valid = si + no
    elettori = values.get("elettori") or 0
    votanti = values.get("votanti") or 0
    yes_share = round2(100 * si / valid) if valid else None
    return {
        "name": name,
        "yes": int(round(si)),
        "no": int(round(no)),
        "yesPercentage": yes_share,
        "electorate": int(elettori) if elettori else None,
        "voters": int(votanti) if votanti else None,
        "blankBallots": int(round(values.get("bianche", 0))) if values.get("bianche") else None,
        "turnout": round2(100 * votanti / elettori) if elettori and votanti else None,
    }


def referendum_records(pack: dict) -> list[dict]:
    records = []
    for question in pack["questions"]:
        national = question["national"]
        valid = (national["yes"] or 0) + (national["no"] or 0)
        lists = []
        if valid:
            lists = [
                {
                    "id": "si",
                    "name": "Sì",
                    "votes": national["yes"],
                    "percentage": national["yesPercentage"],
                    "seats": None,
                    "seatsScope": "non-disponibile",
                },
                {
                    "id": "no",
                    "name": "No",
                    "votes": national["no"],
                    "percentage": round2(100 * national["no"] / valid),
                    "seats": None,
                    "seatsScope": "non-disponibile",
                },
            ]
        regions = []
        for region in question["regions"]:
            regional_valid = (region["yes"] or 0) + (region["no"] or 0)
            regions.append(
                {
                    "id": slugify(region["name"]),
                    "name": region["name"],
                    "electorate": region["electorate"],
                    "voters": region["voters"],
                    "blankBallots": region["blankBallots"],
                    "validVotes": regional_valid or None,
                    "turnout": region["turnout"],
                    "lists": [
                        {
                            "id": "si",
                            "name": "Sì",
                            "votes": region["yes"],
                            "percentage": region["yesPercentage"],
                            "seats": None,
                            "seatsScope": "non-disponibile",
                        },
                        {
                            "id": "no",
                            "name": "No",
                            "votes": region["no"],
                            "percentage": round2(100 * region["no"] / regional_valid) if regional_valid else None,
                            "seats": None,
                            "seatsScope": "non-disponibile",
                        },
                    ],
                }
            )
        note = "Sì e No sono le colonne del file, sommate dai comuni senza stimare i mancanti."
        if pack["date"] == "1946-06-02":
            note += " Per il 2 giugno 1946 il Sì è la repubblica e il No la monarchia."
        records.append(
            {
                "id": question["id"],
                "level": "referendum",
                "chamber": "referendum",
                "date": pack["date"],
                "year": pack["year"],
                "title": f"Referendum {pack['date']} — {question['title'][:140]}",
                "territory": "Italia",
                "source": source(pack["url"], "Zip open data dei referendum."),
                "coverageNotes": [note, "Lo scostamento regionale rispetto al Sì nazionale è un input dei scenari regionali."],
                "turnout": national["turnout"],
                "electorate": national["electorate"],
                "voters": national["voters"],
                "blankBallots": national["blankBallots"],
                "validVotes": valid or None,
                "lists": lists,
                "regions": regions,
                "colleges": [],
            }
        )
    return records


def capoluogo_name(row: dict[str, str]) -> str | None:
    raw = fold(pick(row, "COMUNE", "DESCRIZIONECOMUNE"))
    return CAPOLUOGHI.get(raw)


def collect_capoluoghi(path: str, tables: list[tuple[str, list[dict[str, str]]]]) -> dict[str, dict]:
    date = date_from_name(path)
    if not date:
        return {}
    list_rows = []
    turnout_rows = []
    for name, rows in tables:
        if not rows:
            continue
        kind = classify(name, rows)
        if kind in {"comune", "lists", "lists-naz"}:
            list_rows.extend(rows)
        elif kind == "turnout":
            turnout_rows.extend(rows)
    found: dict[str, dict] = {}
    votes: dict[str, dict[str, float]] = defaultdict(dict)
    seats: dict[str, dict[str, float]] = defaultdict(dict)
    for row in list_rows:
        city = capoluogo_name(row)
        name = list_name(row)
        if not city or not name:
            continue
        amount = list_votes(row)
        if amount is None:
            continue
        # VOTILISTA è il totale di lista ripetuto su ogni candidato: si tiene il massimo, non la somma.
        votes[city][name] = max(votes[city].get(name, 0), amount)
        seat = list_seats(row)
        if seat is not None:
            seats[city][name] = max(seats[city].get(name, 0), seat)
    meta: dict[str, dict[str, float]] = {}
    for row in turnout_rows:
        city = capoluogo_name(row)
        if not city:
            continue
        # primo turno se la colonna c'è
        turno = pick(row, "TURNO")
        if turno and turno not in {"1", "01"}:
            continue
        meta[city] = {
            "elettori": pick_num(row, "ELETTORI", "ELETTORITOT") or 0,
            "votanti": pick_num(row, "NUMVOTANTITOTALI", "VOTANTI", "VOTANTITOT") or 0,
            "bianche": pick_num(row, "SKBIANCHE", "SCHEDEBIANCHE") or 0,
        }
    for city, bucket in votes.items():
        block = territory_block(city, bucket, seats.get(city, {}), meta.get(city, {}), "totale")
        block["date"] = date
        block["sourceUrl"] = BASE + path
        found[city] = block
    return found


def regional_slices(record: dict) -> list[dict]:
    slices = []
    for region in record.get("regions") or []:
        if region["name"] not in REGIONS.values():
            continue
        if not region["lists"]:
            continue
        slices.append(
            {
                "name": region["name"],
                "date": record["date"],
                "sourceUrl": record["source"]["url"],
                "turnout": region["turnout"],
                "lists": region["lists"],
            }
        )
    return slices


def main() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    files = catalog_files()
    selected = []
    for path, name in files:
        if path.startswith("comunali/"):
            date = date_from_name(name) or "1900-01-01"
            if date < "2009-01-01":
                continue
        if path.startswith("camera/") and "20180304" not in path and "20220925" not in path:
            continue
        selected.append(path)
    print("file da leggere", len(selected))

    def download_one(path: str) -> tuple[str, str | None]:
        dest = RAW / path.replace("/", "__")
        try:
            fetch(BASE + path, dest)
            return path, None
        except Exception as error:  # noqa: BLE001
            return path, str(error)

    errors = []
    with ThreadPoolExecutor(max_workers=6) as pool:
        futures = [pool.submit(download_one, path) for path in selected]
        for future in as_completed(futures):
            path, error = future.result()
            if error:
                errors.append({"path": path, "error": error})
                print("ERRORE", path, error)
            else:
                print("ok", path)

    records: dict[str, dict] = {}
    referendum_compact = []
    targets = {}
    latest_region: dict[str, dict] = {}
    latest_city: dict[str, dict] = {}
    failed_parse = []

    for path in sorted(selected):
        dest = RAW / path.replace("/", "__")
        if not dest.exists():
            continue
        family = path.split("/")[0]
        try:
            tables = zip_tables(dest)
        except zipfile.BadZipFile as error:
            failed_parse.append({"path": path, "error": str(error)})
            continue
        if family == "comunali":
            for city, block in collect_capoluoghi(path, tables).items():
                prev = latest_city.get(city)
                if not prev or block["date"] >= prev["date"]:
                    latest_city[city] = block
            continue
        if family == "referendum":
            pack = build_referendum(date_from_name(path) or "", int((date_from_name(path) or "1900")[:4]), BASE + path, [(n, r) for n, r in tables if r])
            if not pack:
                failed_parse.append({"path": path, "error": "nessun quesito"})
                continue
            for record in referendum_records(pack):
                records[record["id"]] = record
            for question in pack["questions"]:
                national_yes = question["national"]["yesPercentage"]
                regions = {
                    region["name"]: region["yesPercentage"]
                    for region in question["regions"]
                    if region["yesPercentage"] is not None
                }
                referendum_compact.append(
                    {
                        "id": question["id"],
                        "date": pack["date"],
                        "title": question["title"],
                        "nationalYes": national_yes,
                        "regions": regions,
                    }
                )
            continue
        record = build_record(family if family != "camera" else "camera", path, tables)
        if not record:
            failed_parse.append({"path": path, "error": "nessuna lista"})
            print("vuoto", path)
            continue
        print(record["id"], "liste", len(record["lists"]), "territori", len(record["regions"]), "affluenza", record["turnout"])
        if family == "camera":
            targets[record["id"]] = {
                "date": record["date"],
                "lists": record["lists"],
                "validVotes": record["validVotes"],
                "source": record["source"],
            }
            continue
        if family == "senato" and record["year"] == 2022:
            continue
        if family == "europee" and record["year"] == 2024:
            targets["europee-2024-zip"] = {
                "date": record["date"],
                "lists": record["lists"],
                "validVotes": record["validVotes"],
                "source": record["source"],
            }
            continue
        if family == "regionali":
            for sl in regional_slices(record):
                prev = latest_region.get(sl["name"])
                if not prev or sl["date"] >= prev["date"]:
                    latest_region[sl["name"]] = sl
            continue
        records[record["id"]] = record
        if family == "europee":
            targets[record["id"]] = {
                "date": record["date"],
                "lists": record["lists"],
                "validVotes": record["validVotes"],
                "source": record["source"],
            }

    # schede archivio per l'ultimo ciclo regionale e per i capoluoghi
    for name, sl in latest_region.items():
        ident = f"regione-{slugify(name)}-{sl['date'][:4]}"
        record = {
            "id": ident,
            "level": "regionali",
            "chamber": "consiglio",
            "date": sl["date"],
            "year": int(sl["date"][:4]),
            "title": f"Regionali {sl['date']} — {name}",
            "territory": name,
            "source": source(sl["sourceUrl"], "Ultimo ciclo regionale trovato negli zip open data per questa regione."),
            "coverageNotes": [
                "È l'ultimo scrutinio di questa regione presente negli zip del catalogo, non l'intera serie dal 1970.",
                "La legge regionale non è codificata: questi voti sono l'input territoriale del simulatore.",
            ],
            "turnout": sl["turnout"],
            "electorate": None,
            "voters": None,
            "blankBallots": None,
            "validVotes": int(sum(item["votes"] or 0 for item in sl["lists"])) or None,
            "lists": sl["lists"],
            "regions": [],
            "colleges": [],
        }
        records[ident] = record
    city_records = []
    for name, block in sorted(latest_city.items()):
        ident = f"capoluogo-{slugify(name)}-{block['date'][:4]}"
        record = {
            "id": ident,
            "level": "comunali",
            "chamber": "comunale",
            "date": block["date"],
            "year": int(block["date"][:4]),
            "title": f"Comunali {block['date']} — {name}",
            "territory": name,
            "source": source(block["sourceUrl"], "Ultimo turno del capoluogo trovato negli zip comunali dal 2009."),
            "coverageNotes": [
                "Capoluogo di provincia o città metropolitana. È il turno più recente nel catalogo dal 2009, di norma il primo turno.",
                "Se la colonna seggi manca, la dimensione del consiglio nel simulatore resta un parametro.",
            ],
            "turnout": block["turnout"],
            "electorate": block["electorate"],
            "voters": block["voters"],
            "blankBallots": block["blankBallots"],
            "validVotes": block["validVotes"],
            "lists": block["lists"],
            "regions": [],
            "colleges": [],
            "municipalities": [block],
        }
        records[ident] = record
        city_records.append(
            {
                "id": ident,
                "name": name,
                "date": block["date"],
                "turnout": block["turnout"],
                "lists": [
                    {"name": item["name"], "percentage": item["percentage"], "votes": item["votes"], "seats": item["seats"]}
                    for item in block["lists"]
                ],
            }
        )

    historical = {
        "generatedAt": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "records": records,
    }
    (OUT / "eligendo-historical.json").write_text(json.dumps(historical, ensure_ascii=False))
    regions_out = {
        "generatedAt": historical["generatedAt"],
        "regions": [
            {
                "name": name,
                "date": sl["date"],
                "turnout": sl["turnout"],
                "sourceUrl": sl["sourceUrl"],
                "lists": [{"name": item["name"], "percentage": item["percentage"], "votes": item["votes"]} for item in sl["lists"]],
            }
            for name, sl in sorted(latest_region.items())
        ],
    }
    (OUT / "eligendo-regions.json").write_text(json.dumps(regions_out, ensure_ascii=False))
    (OUT / "eligendo-capoluoghi.json").write_text(
        json.dumps({"generatedAt": historical["generatedAt"], "cities": city_records}, ensure_ascii=False)
    )
    (OUT / "eligendo-referendum.json").write_text(
        json.dumps({"generatedAt": historical["generatedAt"], "questions": referendum_compact}, ensure_ascii=False)
    )
    (OUT / "eligendo-targets.json").write_text(json.dumps(targets, ensure_ascii=False))

    catalog_path = OUT / "catalog.json"
    catalog = json.loads(catalog_path.read_text())
    kept = []
    for item in catalog["elections"]:
        ident = item["id"]
        if ident.startswith(("politiche-senato-", "europee-", "regione-", "capoluogo-", "referendum-")) and ident not in {
            "politiche-senato-2022",
            "europee-2024",
        }:
            if ident.startswith("politiche-senato-") and ident != "politiche-senato-2022":
                continue
            if ident.startswith("europee-") and ident != "europee-2024":
                continue
            if ident.startswith(("regione-", "capoluogo-", "referendum-")):
                continue
        kept.append(item)
    kept.extend(summary_of(record) for record in records.values())
    kept.sort(key=lambda item: ((item.get("date") or ""), item["id"]))
    catalog["elections"] = kept
    catalog["generatedAt"] = historical["generatedAt"]
    missing_regions = [name for name in REGIONS.values() if name not in latest_region]
    # REGIONS values include duplicates
    missing_regions = sorted(set(missing_regions))
    missing_cities = sorted(set(CAPOLUOGHI.values()) - set(latest_city))
    gaps = []
    for gap in catalog["gaps"]:
        if gap["level"] in {"nazionali", "europee", "regionali", "comunali"}:
            continue
        gaps.append(gap)
    gaps.extend(
        [
            {
                "level": "nazionali",
                "detail": "Camera e Costituente 1946–2022 restano il bollettino Zenodo, con i voti assoluti vuoti. Il Senato 1948–2018 è la somma degli zip open data Eligendo, nazionale e per regione. Camera e Senato 2022 restano i CSV per comune. L'Estero non è nei totali Italia.",
            },
            {
                "level": "europee",
                "detail": "Le europee 1979–2019 sono i totali di lista degli zip open data, con il dettaglio regionale quando il file ha la regione. Il 2024 resta il CSV comunale. I file preferenze non sono stati sommati come voti.",
            },
            {
                "level": "regionali",
                "detail": "Per ogni regione è conservato l'ultimo scrutinio trovato negli zip regionali"
                + (f". Regioni senza uno scrutinio mappato: {', '.join(missing_regions)}." if missing_regions else "."),
            },
            {
                "level": "comunali",
                "detail": f"Capoluoghi con un turno dal 2009: {len(latest_city)}. "
                + (
                    f"Senza un turno in quei file: {', '.join(missing_cities)}. "
                    if missing_cities
                    else "Tutti i capoluoghi cercati hanno almeno un turno. "
                )
                + "Corte Palasio resta il CSV del 28 luglio 2024. Non è l'universo dei comuni.",
            },
            {
                "level": "referendum",
                "detail": f"Referendum nazionali nel catalogo open data: {len(referendum_compact)} quesiti, dal 1946. Il Sì regionale meno il Sì nazionale è il prior territoriale dei scenari regionali. I comuni non mappati su una regione restano nel totale Italia.",
            },
        ]
    )
    catalog["gaps"] = gaps
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2))
    print("record", len(records), "regioni", len(latest_region), "capoluoghi", len(latest_city), "quesiti", len(referendum_compact))
    print("errori download", len(errors), "parse", len(failed_parse))
    (OUT / "eligendo-coverage.json").write_text(
        json.dumps(
            {
                "missingRegions": missing_regions,
                "missingCities": missing_cities,
                "downloadErrors": errors,
                "parseErrors": failed_parse,
                "questionCount": len(referendum_compact),
                "cityCount": len(latest_city),
                "regionCount": len(latest_region),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
